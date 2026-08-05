import csv
import os
from typing import List
from src.ingestion.base import BaseDocumentLoader
from src.metadata.schema import Document, ChunkMetadata
from src.metadata.extractor import MetadataExtractor


class SpreadsheetDocumentLoader(BaseDocumentLoader):
    """Loader for CSV and Excel (.csv, .xlsx, .xls) files."""

    def load(self, source: str) -> str:
        if not os.path.exists(source):
            raise FileNotFoundError(f"Spreadsheet file not found: {source}")

        ext = os.path.splitext(source)[1].lower()
        if ext == ".csv":
            return self._load_csv(source)
        if ext == ".xlsx":
            return self._load_xlsx(source)
        if ext == ".xls":
            return self._load_xls(source)
        raise ValueError(f"Unsupported spreadsheet extension: {ext}")

    def _load_csv(self, source: str) -> str:
        for encoding in ("utf-8-sig", "utf-8", "latin-1", "cp1252"):
            try:
                with open(source, newline="", encoding=encoding) as f:
                    reader = csv.reader(f)
                    return self._rows_to_text("CSV Data", reader)
            except UnicodeDecodeError:
                continue
        raise ValueError(f"Could not decode CSV file: {source}")

    def _load_xlsx(self, source: str) -> str:
        from openpyxl import load_workbook

        wb = load_workbook(source, read_only=True, data_only=True)
        sections: List[str] = []
        try:
            for sheet in wb.worksheets:
                rows = []
                for row in sheet.iter_rows(values_only=True):
                    rows.append([self._cell_str(c) for c in row])
                if rows:
                    sections.append(self._rows_to_text(sheet.title, rows))
        finally:
            wb.close()
        return self._join_sections(sections, source)

    def _load_xls(self, source: str) -> str:
        import xlrd

        book = xlrd.open_workbook(source)
        sections: List[str] = []
        for sheet in book.sheets():
            rows = []
            for rx in range(sheet.nrows):
                rows.append([self._cell_str(sheet.cell_value(rx, cx)) for cx in range(sheet.ncols)])
            if rows:
                sections.append(self._rows_to_text(sheet.name, rows))
        return self._join_sections(sections, source)

    @staticmethod
    def _cell_str(value) -> str:
        if value is None:
            return ""
        return str(value).strip()

    def _rows_to_text(self, title: str, rows) -> str:
        lines = [f"=== {title} ==="]
        for row in rows:
            cells = [c for c in row if c]
            if cells:
                lines.append(" | ".join(cells))
        return "\n".join(lines)

    def _join_sections(self, sections: List[str], source: str) -> str:
        text = "\n\n".join(s for s in sections if s.strip())
        if not text.strip():
            raise ValueError(f"No readable data found in spreadsheet: {source}")
        return text

    def parse(self, raw_content: str, source: str) -> List[Document]:
        metadata = self.extract_metadata(source, raw_content)
        return [Document(content=raw_content.strip(), metadata=metadata)]

    def extract_metadata(self, source: str, raw_content: str) -> ChunkMetadata:
        metadata = MetadataExtractor.extract_from_file(source)
        ext = os.path.splitext(source)[1].lower().lstrip(".")
        metadata.tags.append(ext)
        return metadata
