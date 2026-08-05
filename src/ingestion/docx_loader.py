import os
from typing import List
from src.ingestion.base import BaseDocumentLoader
from src.metadata.schema import Document, ChunkMetadata
from src.metadata.extractor import MetadataExtractor


class DocxDocumentLoader(BaseDocumentLoader):
    """Loader for Microsoft Word (.docx) documents."""

    def load(self, source: str) -> str:
        if not os.path.exists(source):
            raise FileNotFoundError(f"DOCX file not found: {source}")

        try:
            import docx
            doc = docx.Document(source)
            paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
            return "\n\n".join(paragraphs)
        except ImportError:
            with open(source, "r", encoding="utf-8", errors="ignore") as f:
                return f.read()

    def parse(self, raw_content: str, source: str) -> List[Document]:
        metadata = self.extract_metadata(source, raw_content)
        return [Document(content=raw_content.strip(), metadata=metadata)]

    def extract_metadata(self, source: str, raw_content: str) -> ChunkMetadata:
        metadata = MetadataExtractor.extract_from_file(source)
        metadata.tags.append("docx")
        return metadata
